from __future__ import annotations

import argparse
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

APP_NAME = "Унификация Excel-таблиц"
APP_VERSION = "3.1"
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
DEFAULT_OUTPUT_DIR = "Результат"

# Поддерживаемый канонический ряд размеров.
CANONICAL_SIZES: list[Any] = [
    "4XS", "3XS", "2XS", "XS", "S", "M", "L", "XL", "2XL", "3XL",
    "4XL", "5XL", "6XL", "7XL", "8XL", "9XL", "10XL", "S-2XL",
    "164-84-92", "164-96-100",
    "34-5XS", "36-4XS", "38-3XS", "40-2XS", "42-XS", "44-S", "46-M",
    "48-L", "50-XL", "52-2XL", "54-3XL", "56-4XL", "58-5XL", "60-6XL",
    "38-4XS", "40-3XS", "42-2XS", "44-XS", "46-S", "48-M", "50-L",
    "52-XL", "54-2XL", "56-3XL", "58-4XL", "60-5XL", "62-6XL",
    "64-7XL", "66-8XL", "68-9XL", "70-10XL",
    36, 38, 40, 42, 44, 46, 48, 50, 52, 54, 56, 58, 60, 62, 64, 66, 68, 70,
    "3XS-38", "2XS-40", "XS-42", "S-44", "M-46", "L-48", "XL-50",
    "2XL-52", "3XL-54", "4XL-56", "5XL-58", "6XL-60", "7XL-62",
    "8XL-64", "9XL-66", "10XL-68",
    "3XS-36", "2XS-38", "XS-40", "S-42", "M-44", "L-46", "XL-48",
    "2XL-50", "3XL-52", "4XL-54", "5XL-56", "6XL-58", "7XL-60",
]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper().replace("Ё", "Е")
    text = text.replace("Х", "X")  # кириллическая Х
    text = text.replace("–", "-").replace("—", "-").replace("−", "-")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_size(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        number = float(value)
        return str(int(number)) if number.is_integer() else format(number, "g")

    text = normalize_text(value)
    text = text.replace("/", "-").replace("\\", "-")
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"-+", "-", text)

    repeated = re.fullmatch(r"(X+)(S|L)", text)
    if repeated:
        count = len(repeated.group(1))
        suffix = repeated.group(2)
        return f"X{suffix}" if count == 1 else f"{count}X{suffix}"

    # Встречается S-XXL вместо S-2XL.
    text = re.sub(r"XXL", "2XL", text)
    text = re.sub(r"XXXL", "3XL", text)
    text = re.sub(r"XXXXL", "4XL", text)
    text = re.sub(r"XXS", "2XS", text)
    text = re.sub(r"XXXS", "3XS", text)
    text = re.sub(r"XXXXS", "4XS", text)
    return text


SIZE_INDEX = {normalize_size(size): index for index, size in enumerate(CANONICAL_SIZES)}


def header_key(value: Any) -> str:
    text = normalize_text(value)
    return re.sub(r"[^A-ZА-Я0-9]+", "", text)


def as_number(value: Any) -> float | int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return int(value) if float(value).is_integer() else float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text or text.startswith("="):
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def clean_article(value: Any, fallback_text: str = "") -> str:
    """Extract an article, prioritising Modjimoda codes that start with MJM."""
    direct = "" if value is None else str(value).strip()
    combined = "\n".join(part for part in (direct, fallback_text) if part)
    upper = normalize_text(combined)

    # Modjimoda articles are the most reliable identifier in heterogeneous
    # supplier files. They may be in ARTICLE/REF NO. or embedded in STYLE /
    # DESCRIPTION text.
    mjm_match = re.search(r"\b(MJM[A-ZА-Я0-9._/-]{3,})\b", upper, flags=re.IGNORECASE)
    if mjm_match:
        return mjm_match.group(1).strip(" ,.;:")

    labelled_patterns = [
        r"(?:№\s*)?(?:КОД\s*)?АРТИКУЛ[А-ЯA-Z]*\s*[:№,.-]*\s*([A-ZА-Я0-9][A-ZА-Я0-9._/-]{3,})",
        r"(?:ARTICLE|ART\.?|STYLE|ITEM|MODEL|REF)(?:\s*(?:NO|NUMBER|CODE|№))?\s*[:#.,-]*\s*([A-ZА-Я0-9][A-ZА-Я0-9._/-]{3,})",
    ]
    for pattern in labelled_patterns:
        match = re.search(pattern, upper, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip(" ,.;:")

    # A dedicated article cell can contain a code without a label. Reject
    # sentences and service rows such as TOTAL AMOUNT.
    compact = re.sub(r"\s+", "", direct)
    compact_upper = normalize_text(compact)
    rejected = {
        "TOTAL", "TOTALAMOUNT", "ИТОГО", "ОБЩЕЕ", "DEPOSIT", "BALANCE",
        "AMOUNT", "QTY", "QUANTITY", "DESCRIPTION", "STYLE", "ARTICLE",
    }
    if (
        compact
        and not compact.startswith("=")
        and 4 <= len(compact) <= 80
        and re.fullmatch(r"[A-ZА-Я0-9._/-]+", compact_upper)
        and compact_upper not in rejected
        and any(character.isdigit() for character in compact_upper)
    ):
        return compact
    return ""


def extract_dimensions(text: str) -> list[str]:
    found: list[str] = []
    for left, right in re.findall(r"(?<!\d)(\d{2,3})\s*[XХ×*]\s*(\d{2,3})(?!\d)", text, flags=re.IGNORECASE):
        value = f"{int(left)}×{int(right)}"
        if value not in found:
            found.append(value)
    return found


def extract_declared_size(text: str) -> str:
    if not text:
        return ""

    dimensions = extract_dimensions(text)
    if dimensions:
        return "; ".join(dimensions)

    normalized = normalize_text(text)
    match = re.search(
        r"РАЗМЕР(?:Ы)?\s*[:№-]?\s*\(?\s*([^\n\r.;]+?)\s*\)?"
        r"(?=\s*(?:СОСТАВ|ЦВЕТ|БРЕНД|ТНВЭД|КОД|АРТИКУЛ|$))",
        normalized,
        flags=re.IGNORECASE,
    )
    if match:
        candidate = match.group(1).strip(" ():,.-")
        candidate = re.sub(r"\s+", " ", candidate)
        if candidate:
            return candidate

    # Короткий вариант в скобках, например (XS-3XL) или (S).
    for candidate in re.findall(r"\(([^()]{1,30})\)", normalized):
        compact = normalize_size(candidate)
        if compact in SIZE_INDEX or re.fullmatch(r"(?:\d{1,2})?(?:XS|S|M|L|XL)(?:-(?:\d{1,2})?(?:XS|S|M|L|XL))?", compact):
            return candidate.strip()
    return ""


def extract_color(text: str) -> str:
    if not text:
        return ""

    match = re.search(
        r"Цвет\s*[:№-]?\s*([^\n\r.;]+?)"
        r"(?=\s*(?:Состав|Размер|Бренд|ТНВЭД|Код|Артикул|Плотность|Импортер|$))",
        text,
        flags=re.IGNORECASE,
    )
    if match:
        candidate = match.group(1).strip(" ():,.-")
        candidate = re.sub(r"\s+", " ", candidate)
        if candidate:
            return candidate
    return ""


def single_canonical_size(size_text: str) -> str | None:
    if not size_text:
        return None
    normalized = normalize_size(size_text)
    if normalized in SIZE_INDEX:
        return normalized
    return None


def size_columns_in_row(row: Sequence[Any]) -> dict[int, str]:
    return {
        column_index: normalized
        for column_index, value in enumerate(row)
        if (normalized := normalize_size(value)) in SIZE_INDEX
    }


@dataclass
class SheetData:
    name: str
    rows: list[list[Any]]


@dataclass
class Candidate:
    sheet: SheetData
    header_row: int
    data_start_row: int
    article_col: int | None
    name_col: int
    quantity_col: int | None
    total_col: int | None
    explicit_size_col: int | None
    color_col: int | None
    size_columns: dict[int, str]
    packing_mode: bool
    score: int


@dataclass
class ProductRow:
    name: str
    article: str
    color: str = ""
    declared_size: str = ""
    sizes: dict[str, float | int] = field(default_factory=dict)
    total: float | int | None = None


def find_libreoffice() -> str | None:
    for executable in ("libreoffice", "soffice"):
        path = shutil.which(executable)
        if path:
            return path
    return None


def read_xlsx(path: Path) -> list[SheetData]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: list[SheetData] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sheets.append(SheetData(worksheet.title, rows))
    finally:
        workbook.close()
    return sheets


def read_xls_with_xlrd(path: Path) -> list[SheetData]:
    import xlrd  # optional dependency for legacy .xls

    book = xlrd.open_workbook(str(path), on_demand=True)
    sheets: list[SheetData] = []
    try:
        for sheet_name in book.sheet_names():
            sheet = book.sheet_by_name(sheet_name)
            rows = [sheet.row_values(index) for index in range(sheet.nrows)]
            sheets.append(SheetData(sheet_name, rows))
    finally:
        book.release_resources()
    return sheets


def read_xls(path: Path) -> list[SheetData]:
    try:
        return read_xls_with_xlrd(path)
    except ImportError:
        office = find_libreoffice()
        if not office:
            raise RuntimeError(
                "Для старого формата .xls установите пакет xlrd "
                "(python -m pip install xlrd) или LibreOffice."
            )
        with tempfile.TemporaryDirectory(prefix="excel_unifier_") as temp_dir:
            temp_path = Path(temp_dir)
            profile_dir = temp_path / "libreoffice_profile"
            runtime_dir = temp_path / "runtime"
            profile_dir.mkdir(parents=True, exist_ok=True)
            runtime_dir.mkdir(parents=True, exist_ok=True)
            try:
                runtime_dir.chmod(0o700)
            except OSError:
                pass

            environment = os.environ.copy()
            environment["HOME"] = temp_dir
            environment["XDG_RUNTIME_DIR"] = str(runtime_dir)
            result = subprocess.run(
                [
                    office,
                    "--headless",
                    f"-env:UserInstallation={profile_dir.as_uri()}",
                    "--convert-to", "xlsx",
                    "--outdir", temp_dir,
                    str(path),
                ],
                capture_output=True,
                text=True,
                timeout=180,
                env=environment,
            )
            converted = temp_path / f"{path.stem}.xlsx"
            if result.returncode != 0 or not converted.exists():
                details = (result.stderr or result.stdout or "неизвестная ошибка").strip()
                raise RuntimeError(f"LibreOffice не смог преобразовать .xls: {details}")
            return read_xlsx(converted)


def read_workbook(path: Path) -> list[SheetData]:
    extension = path.suffix.lower()
    if extension in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    if extension == ".xls":
        return read_xls(path)
    raise ValueError(f"Неподдерживаемое расширение: {extension}")


def choose_column(headers: Sequence[Any], predicates: Sequence[tuple[int, callable]]) -> int | None:
    best: tuple[int, int] | None = None
    for index, value in enumerate(headers):
        key = header_key(value)
        text = normalize_text(value)
        for weight, predicate in predicates:
            if predicate(key, text):
                candidate = (weight, -index)
                if best is None or candidate > best:
                    best = candidate
                    best_index = index
    return None if best is None else best_index


def row_value(row: Sequence[Any], column: int | None) -> Any:
    if column is None or column < 0 or column >= len(row):
        return None
    return row[column]


def row_text(row: Sequence[Any]) -> str:
    return "\n".join(
        str(value).strip()
        for value in row
        if value not in (None, "")
    )


def is_service_or_total_row(row: Sequence[Any]) -> bool:
    first_values = [
        normalize_text(value)
        for value in row[:6]
        if value not in (None, "")
    ]
    if not first_values:
        return False
    first = first_values[0]
    service_prefixes = (
        "TOTAL", "GRAND TOTAL", "ИТОГО", "ОБЩЕЕ", "DEPOSIT", "THE BALANCE",
        "BALANCE", "SAY TOTAL", "BANK DETAIL", "BENEFICIARY",
        "SWIFT", "AUTHORIZED SIGNATURE", "E.&O.E",
    )
    if any(first.startswith(prefix) for prefix in service_prefixes):
        return True

    # Pivot/summary sheets commonly put subtotals after the model name, for
    # example "Брюки трикотажные мужские Total".  Looking only at the start of
    # the cell lets such a row inherit the preceding article and doubles the
    # quantities when both detail rows and their subtotal are parsed.
    total_suffix = re.compile(r"(?:^|[\s:;,.()/_-])(GRAND\s+TOTAL|SUB[ -]?TOTAL|TOTAL|ИТОГО|ВСЕГО)$")
    return any(total_suffix.search(value) for value in first_values)


def summary_header_score(headers: Sequence[Any], next_header: Sequence[Any]) -> int:
    """Recognise a product summary by its columns, independently of sheet name."""
    keys = {header_key(value) for value in [*headers, *next_header] if value not in (None, "")}

    has_photo = any(
        key.startswith(("ФОТО", "PHOTO", "IMAGE", "PICTURE"))
        for key in keys
    )
    has_model_name = any(
        key in {"НАЗВАНИЕМОДЕЛИ", "НАИМЕНОВАНИЕМОДЕЛИ", "MODELNAME", "PRODUCTNAME"}
        or key.startswith(("НАЗВАНИЕМОДЕЛИ", "НАИМЕНОВАНИЕМОДЕЛИ"))
        for key in keys
    )
    has_article = any(
        "АРТИКУЛ" in key
        or key.startswith(("ARTICLE", "ARTNO", "SKU", "STYLECODE", "MODELNO"))
        for key in keys
    )
    has_color = any(
        key in {"ЦВЕТ", "COLOR", "COLOUR"}
        or key.startswith(("НАЗВАНИЕЦВЕТА", "НАИМЕНОВАНИЕЦВЕТА"))
        for key in keys
    )

    # The four-column signature is deliberately strong: it identifies the
    # user-facing product summary even if the sender renames the worksheet.
    if has_photo and has_model_name and has_article and has_color:
        return 1_000
    if has_photo and has_model_name and has_article:
        return 500
    return 0


def clean_product_name(value: Any, article: str = "") -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip()
    if not text:
        return ""

    # Supplier descriptions often put the concise product name on the first
    # line and then add composition, trademark and manufacturer metadata.
    lines = [re.sub(r"\s+", " ", line).strip(" ,;") for line in text.splitlines()]
    useful: list[str] = []
    stop_prefixes = (
        "COMPOSITION", "СОСТАВ", "TRADEMARK", "ТОРГОВАЯ МАРКА",
        "MANUFACTURER", "ПРОИЗВОДИТЕЛЬ", "COUNTRY OF ORIGIN",
        "СТРАНА ПРОИСХОЖДЕНИЯ", "ARTICLE", "АРТИКУЛ",
    )
    for line in lines:
        if not line:
            continue
        normalized = normalize_text(line)
        if article and normalize_text(article) == normalized:
            continue
        if any(normalized.startswith(prefix) for prefix in stop_prefixes):
            break
        if re.fullmatch(r"MJM[A-ZА-Я0-9._/-]{3,}", normalized):
            continue
        useful.append(line)

    if useful:
        return " ".join(useful)
    return article or re.sub(r"\s+", " ", text)


def find_article_content_column(
    sheet: SheetData,
    start_row: int,
    max_rows: int = 120,
) -> int | None:
    counts: dict[int, int] = {}
    for row in sheet.rows[start_row:start_row + max_rows]:
        for column_index, value in enumerate(row):
            if value in (None, ""):
                continue
            if re.search(r"\bMJM[A-ZА-Я0-9._/-]{3,}\b", normalize_text(value)):
                counts[column_index] = counts.get(column_index, 0) + 1
    if not counts:
        return None
    return max(counts, key=lambda index: (counts[index], -index))


def detect_candidates(sheets: Sequence[SheetData]) -> list[Candidate]:
    candidates: list[Candidate] = []

    article_rules = [
        (120, lambda key, text: "АРТИКУЛ" in key and "ЦВЕТ" not in key),
        (115, lambda key, text: key in {
            "SKU", "ARTICLE", "ARTICLENO", "ARTICLENUMBER", "ARTICLECODE",
            "ARTNO", "ARTNUMBER", "REFNO", "REFERENCENO", "ITEMNO",
            "ITEMCODE", "STYLECODE", "STYLENO", "MODELNO", "КОДАРТИКУЛА",
        }),
        (100, lambda key, text: key.startswith("ARTICLE") or key.startswith("ARTNO")),
        (90, lambda key, text: key.startswith("REFNO") or key.startswith("ITEMNO")),
    ]
    name_rules = [
        (110, lambda key, text: "ОПИСАНИЕТОВАРА" in key or "ОПИСАНИЕТOВАРА" in key),
        (105, lambda key, text: "НАИМЕНОВАНИЕМОДЕЛИ" in key),
        (100, lambda key, text: "НАЗВАНИЕИЗДЕЛИЯ" in key),
        (95, lambda key, text: "НАИМЕНОВАНИЕ" in key or "НАЗВАНИЕМОДЕЛИ" in key),
        (92, lambda key, text: key in {
            "DESCRIPTION", "ITEMDESCRIPTION", "PRODUCTDESCRIPTION",
            "GOODSDESCRIPTION", "PRODUCTNAME", "ITEMNAME",
        }),
        (88, lambda key, text: key.startswith("DESCRIPTION")),
        (82, lambda key, text: key == "STYLE" or key.startswith("STYLEDESCRIPTION")),
        (75, lambda key, text: "ОПИСАНИЕ" in key),
    ]
    quantity_rules = [
        (120, lambda key, text: key in {"TOTALQTY", "TOTALQUANTITY", "QTY", "QUANTITY", "QUANTITIES"}),
        (115, lambda key, text: key.startswith("TOTALQTY") or key.startswith("TOTALQUANTITY")),
        (110, lambda key, text: "КОЛВОИЗДЕЛИЙ" in key or "КОЛВОШТУК" in key),
        (105, lambda key, text: "КОЛВОЕДИНИЦ" in key or "КОЛВОЕДЕНИЦ" in key),
        (100, lambda key, text: key in {"КОЛВО", "КОЛИЧЕСТВО"}),
        (90, lambda key, text: key.startswith("QUANTIT")),
        (65, lambda key, text: "КОЛВОКОМПЛЕКТОВ" in key),
        (40, lambda key, text: "КОЛВО" in key and "МЕШ" not in key and "КГ" not in key),
    ]
    total_rules = [
        (110, lambda key, text: key.startswith("ОБЩЕЕ") or key.startswith("ИТОГО")),
        (95, lambda key, text: key in {"TOTAL", "TOTALPCS"}),
    ]
    size_text_rules = [
        (110, lambda key, text: "ФАКТИЧЕСКИЙРАЗМЕР" in key),
        (100, lambda key, text: key in {"РАЗМЕР", "РАЗМЕРЫ", "SIZE", "SIZES"}),
    ]
    color_rules = [
        (110, lambda key, text: key in {"ЦВЕТ", "COLOR", "COLOUR"}),
        (105, lambda key, text: "НАЗВАНИЕЦВЕТА" in key or "НАИМЕНОВАНИЕЦВЕТА" in key),
        (90, lambda key, text: "ЦВЕТ" in key and "КОД" not in key and "АРТИКУЛ" not in key),
        (70, lambda key, text: key in {"КОДЦВЕТА", "COLORCODE", "COLOURCODE"}),
    ]

    for sheet in sheets:
        max_scan = min(len(sheet.rows), 120)
        for row_index in range(max_scan):
            headers = sheet.rows[row_index]
            if not any(value not in (None, "") for value in headers):
                continue

            next_header = sheet.rows[row_index + 1] if row_index + 1 < len(sheet.rows) else []
            field_headers = list(headers)
            if len(next_header) > len(field_headers):
                field_headers.extend([None] * (len(next_header) - len(field_headers)))
            for index, value in enumerate(next_header):
                if index >= len(field_headers):
                    field_headers.append(value)
                elif field_headers[index] in (None, ""):
                    field_headers[index] = value

            article_col = choose_column(headers, article_rules)
            name_col = choose_column(headers, name_rules)
            if name_col is None:
                name_col = choose_column(field_headers, name_rules)

            quantity_col = choose_column(headers, quantity_rules)
            total_col = choose_column(headers, total_rules)
            explicit_size_col = choose_column(headers, size_text_rules)
            color_col = choose_column(headers, color_rules)

            if quantity_col is None:
                quantity_col = choose_column(field_headers, quantity_rules)
            if total_col is None:
                total_col = choose_column(field_headers, total_rules)
            if explicit_size_col is None:
                explicit_size_col = choose_column(field_headers, size_text_rules)
            if color_col is None:
                color_col = choose_column(field_headers, color_rules)

            current_size_columns = size_columns_in_row(headers)
            next_size_columns = size_columns_in_row(next_header)

            # A second header line is considered a size header only when it
            # contains at least two recognised sizes. A single "M" in the next
            # row is usually already the first product.
            if len(next_size_columns) >= 2 and len(next_size_columns) > len(current_size_columns):
                size_columns = next_size_columns
                data_start_row = row_index + 2
            else:
                size_columns = current_size_columns
                data_start_row = row_index + 1

            if name_col is None:
                continue

            content_article_col = find_article_content_column(sheet, data_start_row)
            if article_col is None:
                article_col = content_article_col

            if quantity_col is None and total_col is None and not size_columns:
                continue

            next_rows = sheet.rows[data_start_row:data_start_row + 120]
            article_count = 0
            name_count = 0
            numeric_count = 0
            for row in next_rows:
                article_value = row_value(row, article_col)
                name_value = row_value(row, name_col)
                if clean_article(article_value, row_text(row)):
                    article_count += 1
                if name_value not in (None, ""):
                    name_count += 1
                for column_index in [quantity_col, total_col, *size_columns.keys()]:
                    if as_number(row_value(row, column_index)) is not None:
                        numeric_count += 1
                        break

            if article_count == 0:
                continue

            all_header_keys = {header_key(value) for value in [*headers, *next_header]}
            sheet_name = normalize_text(sheet.name)
            packing_mode = (
                "УПАК" in sheet_name
                or "PACK" in sheet_name
                or sheet_name == "PL"
                or any(
                    key.startswith(prefix)
                    for key in all_header_keys
                    for prefix in ("PCSPKG", "PKGNO", "CTNS", "PACKAGES")
                )
            )

            score = 100 + min(article_count, 30) * 4 + min(name_count, 30)
            score += len(size_columns) * 10
            score += 50 if quantity_col is not None else 0
            score += 35 if total_col is not None else 0
            score += 15 if color_col is not None else 0
            score += min(numeric_count, 20) * 2
            score += 30 if content_article_col is not None else 0
            score += summary_header_score(headers, next_header)

            # Prefer commercial invoices over packing lists when both are in
            # the same workbook: invoice size cells normally contain final
            # item quantities, while a packing list can contain PCS/PKG.
            if (
                "ИНВОЙС" in sheet_name
                or "INVOICE" in sheet_name
                or sheet_name in {"CI", "C/I"}
                or "COMMERCIAL" in sheet_name
            ):
                score += 260
            if "С ФОТО" in sheet_name:
                score += 170
            if "СПЕЦ" in sheet_name:
                score += 40
            if packing_mode:
                score += 60

            candidates.append(
                Candidate(
                    sheet=sheet,
                    header_row=row_index,
                    data_start_row=data_start_row,
                    article_col=article_col,
                    name_col=name_col,
                    quantity_col=quantity_col,
                    total_col=total_col,
                    explicit_size_col=explicit_size_col,
                    color_col=color_col,
                    size_columns=size_columns,
                    packing_mode=packing_mode,
                    score=score,
                )
            )

    return sorted(candidates, key=lambda item: item.score, reverse=True)


def parse_candidate(candidate: Candidate) -> list[ProductRow]:
    rows: list[ProductRow] = []
    empty_after_data = 0
    started = False

    current_article = ""
    current_name = ""
    current_color = ""
    size_columns = candidate.size_columns

    for source_row in candidate.sheet.rows[candidate.data_start_row:]:
        if not any(value not in (None, "") for value in source_row):
            if started:
                empty_after_data += 1
                if empty_after_data >= 25:
                    break
            continue

        full_text = row_text(source_row)
        explicit_name_value = row_value(source_row, candidate.name_col)
        explicit_name_raw = "" if explicit_name_value is None else str(explicit_name_value).strip()
        article_value = row_value(source_row, candidate.article_col)
        found_article = clean_article(article_value, full_text)

        replacement_size_columns = size_columns_in_row(source_row)
        if len(replacement_size_columns) >= 2 and not found_article:
            size_columns = replacement_size_columns
            continue

        # Do not let totals and bank/service sections inherit the preceding
        # product article.  Skip them even if a supplier repeats the article
        # in a subtotal row.
        if is_service_or_total_row(source_row):
            if started:
                empty_after_data += 1
            continue

        if found_article:
            current_article = found_article
            cleaned_name = clean_product_name(explicit_name_raw, found_article)
            if cleaned_name:
                current_name = cleaned_name

            explicit_color = row_value(source_row, candidate.color_col)
            if explicit_color not in (None, ""):
                current_color = re.sub(r"\s+", " ", str(explicit_color).strip())
            elif explicit_name_raw:
                extracted = extract_color(explicit_name_raw)
                if extracted:
                    current_color = extracted

        article = found_article or current_article
        if not article:
            if started:
                empty_after_data += 1
                if empty_after_data >= 25:
                    break
            continue

        started = True
        empty_after_data = 0

        row_name = clean_product_name(explicit_name_raw, article)
        name = row_name or current_name or article
        if row_name:
            current_name = row_name

        color = ""
        explicit_color = row_value(source_row, candidate.color_col)
        if explicit_color not in (None, ""):
            color = re.sub(r"\s+", " ", str(explicit_color).strip())
            current_color = color
        else:
            color = current_color or extract_color(explicit_name_raw)

        declared_size = ""
        explicit_size = row_value(source_row, candidate.explicit_size_col)
        if explicit_size not in (None, ""):
            declared_size = str(explicit_size).strip()
        if not declared_size:
            declared_size = extract_declared_size(explicit_name_raw or name)

        mapped_sizes: dict[str, float | int] = {}
        for column_index, normalized_size in size_columns.items():
            quantity = as_number(row_value(source_row, column_index))
            if quantity is not None and quantity != 0:
                mapped_sizes[normalized_size] = quantity

        total = None
        for column_index in (candidate.total_col, candidate.quantity_col):
            value = as_number(row_value(source_row, column_index))
            if value is not None:
                total = value
                break

        # Wide invoice sheets often end with a totals row where all product
        # identity cells are empty, while several size cells and the total are
        # populated. Such a row must not inherit the preceding MJM article.
        identity_cells_empty = (
            not found_article
            and not explicit_name_raw
            and row_value(source_row, candidate.color_col) in (None, "")
            and row_value(source_row, candidate.explicit_size_col) in (None, "")
        )
        if size_columns and identity_cells_empty and mapped_sizes:
            continue

        if mapped_sizes:
            # In packing lists one marked size column can contain PCS/PKG,
            # while TOTAL QTY contains the real number of pieces. Preserve the
            # size position but use the final quantity.
            if (
                candidate.packing_mode
                and total is not None
                and len(mapped_sizes) == 1
                and float(total) >= float(next(iter(mapped_sizes.values())))
            ):
                only_size = next(iter(mapped_sizes))
                mapped_sizes = {only_size: total}
            else:
                calculated_total = sum(float(value) for value in mapped_sizes.values())
                total = int(calculated_total) if calculated_total.is_integer() else calculated_total
        elif total is not None:
            exact_size = single_canonical_size(declared_size)
            if exact_size:
                mapped_sizes[exact_size] = total

        if total is None:
            continue

        rows.append(
            ProductRow(
                name=name,
                article=article,
                color=color,
                declared_size=declared_size,
                sizes=mapped_sizes,
                total=total,
            )
        )

    return rows


def merge_products(rows: Sequence[ProductRow]) -> list[ProductRow]:
    merged: OrderedDict[tuple[str, str, str], ProductRow] = OrderedDict()
    for row in rows:
        key = (row.article.casefold(), row.name.casefold(), row.color.casefold())
        if key not in merged:
            merged[key] = ProductRow(
                name=row.name,
                article=row.article,
                color=row.color,
                declared_size=row.declared_size,
                sizes=dict(row.sizes),
                total=row.total,
            )
            continue

        current = merged[key]
        if not current.declared_size and row.declared_size:
            current.declared_size = row.declared_size
        elif (
            row.declared_size
            and current.declared_size
            and normalize_size(row.declared_size) != normalize_size(current.declared_size)
        ):
            # For garments represented by several source rows, the actual
            # distribution is already present in the size columns. Avoid
            # leaving a misleading single size in the text field.
            if current.sizes or row.sizes:
                current.declared_size = ""
            else:
                values = [
                    value for value in current.declared_size.split("; ")
                    if value
                ]
                if row.declared_size not in values:
                    values.append(row.declared_size)
                current.declared_size = "; ".join(values)

        for size, quantity in row.sizes.items():
            current.sizes[size] = current.sizes.get(size, 0) + quantity
        if row.total is not None:
            current.total = (current.total or 0) + row.total

    for product in merged.values():
        if len(product.sizes) > 1:
            product.declared_size = ""
    return list(merged.values())


def candidate_result_score(candidate: Candidate, products: Sequence[ProductRow]) -> int:
    """Prefer the candidate that contains the actual size distribution."""
    distinct_articles = len({product.article.casefold() for product in products})
    sized_products = sum(1 for product in products if product.sizes)
    multi_size_products = sum(1 for product in products if len(product.sizes) >= 2)

    return (
        candidate.score
        + min(distinct_articles, 20) * 5
        + min(sized_products, 20) * 20
        + min(multi_size_products, 20) * 120
    )


def extract_products(path: Path) -> tuple[list[ProductRow], Candidate]:
    sheets = read_workbook(path)
    candidates = detect_candidates(sheets)
    if not candidates:
        raise RuntimeError(
            "Не найдена подходящая таблица. Программа искала наименование/описание, количество или размерные колонки и артикул — в том числе код, начинающийся с MJM."
        )

    errors: list[str] = []
    successful: list[tuple[int, list[ProductRow], Candidate]] = []
    for candidate in candidates:
        parsed = parse_candidate(candidate)
        if not parsed:
            errors.append(f'{candidate.sheet.name}: строки товаров не распознаны')
            continue
        merged = merge_products(parsed)
        successful.append((candidate_result_score(candidate, merged), merged, candidate))

    if successful:
        successful.sort(key=lambda item: item[0], reverse=True)
        _, products, candidate = successful[0]
        for product in products:
            matching_articles = {
                reference.article
                for _, references, _ in successful[1:]
                for reference in references
                if (
                    reference.article != product.article
                    and len(reference.article) == len(product.article)
                    and sum(
                        left != right
                        for left, right in zip(
                            reference.article.casefold(),
                            product.article.casefold(),
                        )
                    ) == 1
                    and reference.sizes == product.sizes
                    and normalize_text(reference.color) == normalize_text(product.color)
                    and len(reference.name) > len(product.name)
                )
            }
            if len(matching_articles) == 1:
                product.article = matching_articles.pop()
        return products, candidate

    raise RuntimeError("; ".join(errors[:5]))


def display_size(size_key: str) -> str:
    for canonical in CANONICAL_SIZES:
        if normalize_size(canonical) == size_key:
            return str(canonical)
    return size_key


def expanded_output_rows(products: Sequence[ProductRow]) -> list[list[Any]]:
    """Convert products to the long format: one row per article and size."""
    result: list[list[Any]] = []
    sequence = 1

    for product in products:
        if product.sizes:
            ordered_sizes = sorted(
                product.sizes.items(),
                key=lambda item: (SIZE_INDEX.get(item[0], 10_000), item[0]),
            )
            for size_key, quantity in ordered_sizes:
                if quantity in (None, 0):
                    continue
                result.append([
                    sequence,
                    product.name,
                    product.article,
                    product.color,
                    display_size(size_key),
                    quantity,
                ])
                sequence += 1
            continue

        # Physical towel dimensions are part of the product description, not
        # an apparel variation. The requested output keeps their Size empty.
        output_size = product.declared_size
        if extract_dimensions(output_size):
            output_size = ""

        result.append([
            sequence,
            product.name,
            product.article,
            product.color,
            output_size,
            product.total,
        ])
        sequence += 1

    return result


def build_output(products: Sequence[ProductRow], output_path: Path, source_path: Path) -> int:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Таблица"

    headers = ["№", "Наименование", "Артикул", "Цвет", "Размер", "Количество"]
    worksheet.append(headers)

    output_rows = expanded_output_rows(products)
    for row in output_rows:
        worksheet.append(row)

    last_row = worksheet.max_row
    last_col = worksheet.max_column
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    worksheet.row_dimensions[1].height = 30

    for row in worksheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 2)
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[2].number_format = "@"
        row[3].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row[4].alignment = Alignment(horizontal="center", vertical="center")
        row[5].alignment = Alignment(horizontal="center", vertical="center")
        row[5].number_format = "0"
        name_length = len(str(row[1].value or ""))
        estimated_lines = max(1, math.ceil(name_length / 85))
        worksheet.row_dimensions[row[0].row].height = min(72, 20 * estimated_lines + 4)

    worksheet.column_dimensions["A"].width = 7
    worksheet.column_dimensions["B"].width = 70
    worksheet.column_dimensions["C"].width = 24
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 15
    worksheet.column_dimensions["F"].width = 14

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.print_title_rows = "1:1"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    if last_row >= 2:
        table = Table(displayName="UnifiedProducts", ref=f"A1:F{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    else:
        worksheet.auto_filter.ref = "A1:F1"

    workbook.properties.title = f"Унифицированная таблица: {source_path.name}"
    workbook.properties.subject = "Наименование, артикул, цвет, размер и количество"
    workbook.properties.creator = APP_NAME

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_path.with_name(f".{output_path.name}.tmp.xlsx")
    workbook.save(temp_path)
    os.replace(temp_path, output_path)
    return len(output_rows)

def output_path_for(source: Path, output_dir_name: str) -> Path:
    output_dir = source.parent / output_dir_name
    # openpyxl сохраняет только современный XLSX. Для исходного XLSX имя полностью совпадает.
    filename = source.name if source.suffix.lower() in {".xlsx", ".xlsm"} else f"{source.stem}.xlsx"
    if filename.lower().endswith(".xlsm"):
        filename = f"{Path(filename).stem}.xlsx"
    output_path = output_dir / filename
    if output_path.resolve() == source.resolve():
        raise ValueError("Папка результата не должна перезаписывать исходный файл.")
    return output_path


def convert_file(source: Path, output_dir_name: str = DEFAULT_OUTPUT_DIR) -> tuple[Path, Candidate, int]:
    source = source.expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    if source.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Поддерживаются: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    products, candidate = extract_products(source)
    output_path = output_path_for(source, output_dir_name)
    output_row_count = build_output(products, output_path, source)
    return output_path, candidate, output_row_count


def run_cli(files: Iterable[str], output_dir_name: str) -> int:
    failed = 0
    for file_name in files:
        source = Path(file_name)
        try:
            output, candidate, count = convert_file(source, output_dir_name)
            print(
                f"ГОТОВО: {source.name} -> {output}\n"
                f"  источник данных: лист «{candidate.sheet.name}», строка заголовков {candidate.header_row + 1}\n"
                f"  строк результата: {count}"
            )
        except Exception as error:
            failed += 1
            print(f"ОШИБКА: {source}: {error}", file=sys.stderr)
    return 1 if failed else 0


def run_gui(output_dir_name: str) -> int:
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except ImportError:
        print("Tkinter не установлен. Запустите программу с путями к файлам в командной строке.", file=sys.stderr)
        return 2

    root = tk.Tk()
    root.title(APP_NAME)
    root.geometry("800x520")
    root.minsize(650, 420)

    selected: list[Path] = []
    frame = ttk.Frame(root, padding=16)
    frame.pack(fill="both", expand=True)

    ttk.Label(frame, text=APP_NAME, font=("Segoe UI", 16, "bold")).pack(anchor="w")
    ttk.Label(
        frame,
        text=(
            "Выберите исходные Excel-файлы. Результат в формате «одна строка = один размер» "
            "будет создан в папке "
            f"«{output_dir_name}» рядом с каждым исходником и получит то же имя."
        ),
        wraplength=750,
        justify="left",
    ).pack(anchor="w", pady=(6, 14))

    listbox = tk.Listbox(frame, height=10)
    listbox.pack(fill="both", expand=True)

    log = tk.Text(frame, height=8, state="disabled", wrap="word")
    log.pack(fill="both", expand=True, pady=(12, 0))

    def append_log(text: str) -> None:
        log.configure(state="normal")
        log.insert("end", text + "\n")
        log.see("end")
        log.configure(state="disabled")
        root.update_idletasks()

    def choose_files() -> None:
        paths = filedialog.askopenfilenames(
            title="Выберите Excel-файлы",
            filetypes=[("Excel", "*.xlsx *.xlsm *.xls"), ("Все файлы", "*.*")],
        )
        for value in paths:
            path = Path(value)
            if path not in selected:
                selected.append(path)
                listbox.insert("end", str(path))
        run_button.configure(state="normal" if selected else "disabled")

    def clear_files() -> None:
        selected.clear()
        listbox.delete(0, "end")
        run_button.configure(state="disabled")

    def process() -> None:
        run_button.configure(state="disabled")
        failures: list[str] = []
        for source in selected:
            append_log(f"Обработка: {source.name}")
            try:
                output, candidate, count = convert_file(source, output_dir_name)
                append_log(
                    f"  Готово: {output}\n"
                    f"  Лист: {candidate.sheet.name}; строк результата: {count}"
                )
            except Exception as error:
                failures.append(f"{source.name}: {error}")
                append_log(f"  Ошибка: {error}")
        run_button.configure(state="normal" if selected else "disabled")
        if failures:
            messagebox.showwarning("Завершено с ошибками", "\n".join(failures))
        else:
            messagebox.showinfo("Готово", "Все файлы преобразованы.")

    buttons = ttk.Frame(frame)
    buttons.pack(fill="x", pady=(12, 0))
    ttk.Button(buttons, text="Выбрать файлы", command=choose_files).pack(side="left")
    run_button = ttk.Button(buttons, text="Преобразовать", command=process, state="disabled")
    run_button.pack(side="left", padx=(8, 0))
    ttk.Button(buttons, text="Очистить", command=clear_files).pack(side="left", padx=(8, 0))

    root.mainloop()
    return 0


def self_test() -> None:
    assert normalize_size("XXS") == "2XS"
    assert normalize_size("XXXL") == "3XL"
    assert normalize_size("S - XXL") == "S-2XL"
    assert extract_dimensions("Размер 30 х 60 и 50x80") == ["30×60", "50×80"]
    assert single_canonical_size("S") == "S"
    assert single_canonical_size("XS-3XL") is None
    assert extract_color("Размер: 30х30\nЦвет: Микс\nСостав: Хлопок 100%") == "Микс"
    assert clean_article(None, "DEMO JACKET\nMJMDEMOJACKET001") == "MJMDEMOJACKET001"
    assert clean_article("MJMDEMOJACKET002", "") == "MJMDEMOJACKET002"
    assert clean_article("TOTAL AMOUNT", "") == ""
    assert clean_product_name("DEMO JACKET\nCOMPOSITION: DEMO FABRIC\nMJMDEMOJACKET001", "MJMDEMOJACKET001") == "DEMO JACKET"
    print("Self-test: OK")


def main() -> int:
    parser = argparse.ArgumentParser(description=APP_NAME)
    parser.add_argument("--version", action="version", version=f"%(prog)s {APP_VERSION}")
    parser.add_argument("files", nargs="*", help="Исходные .xlsx/.xlsm/.xls")
    parser.add_argument(
        "--output-dir",
        default=DEFAULT_OUTPUT_DIR,
        help=f"Имя папки результата рядом с исходником (по умолчанию: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()

    if args.self_test:
        self_test()
        return 0
    if args.files:
        return run_cli(args.files, args.output_dir)
    return run_gui(args.output_dir)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        raise SystemExit(130)
    except Exception:
        traceback.print_exc()
        raise SystemExit(1)
