"""Invented fixtures only; no supplier/customer files or external office process."""
import hashlib
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import excel_unifier as unifier


HEADERS = ["№", "Наименование", "Артикул", "Цвет", "Размер", "Количество"]


def parse_sheets(*sheets):
    with patch.object(unifier, "read_workbook", return_value=list(sheets)):
        return unifier.extract_products(Path("synthetic.xlsx"))


def invented_sheets():
    summary = unifier.SheetData("Synthetic summary", [
        ["Наименование", "Артикул", "Цвет", "XS", "S", "M", "L", "Общее"],
        ["Учебные брюки", "MJMDEMO0002P", "Синий", 0, 0, 0, 17, 17],
        ["Учебный костюм", "MJMDEMO0001A", "Зелёный", 3, 5, 7, 11, 26],
        ["Grand Total", "", "", 3, 5, 7, 28, 43],
    ])
    detail = unifier.SheetData("Synthetic detail", [
        ["Фото", "Наименование", "Артикул", "Цвет", "Цена", "XS", "S", "M", "L", "Общее"],
        ["", "Учебный костюм с демонстрационной окантовкой", "MJMDEMO0001B", "Зелёный", 1.75, 3, 5, 7, 11, 26],
    ])
    return summary, detail


class SyntheticParserTests(unittest.TestCase):
    def setUp(self):
        # An accidental legacy conversion must fail closed, never launch office.
        office = patch.object(unifier, "find_libreoffice", return_value=None)
        office.start()
        self.addCleanup(office.stop)

    def test_prefix_priority_and_direct_article_fallback(self):
        self.assertEqual(unifier.clean_article("GENERIC0009", "ARTICLE MJMDEMO0100"), "MJMDEMO0100")
        self.assertEqual(unifier.clean_article("000123", ""), "000123")
        self.assertEqual(unifier.clean_article("TOTAL AMOUNT", ""), "")

    def test_size_normalization_and_invalid_quantities(self):
        self.assertEqual(unifier.normalize_size("XXS"), "2XS")
        self.assertEqual(unifier.normalize_size("XXXL"), "3XL")
        self.assertEqual(unifier.normalize_size("S - XXL"), "S-2XL")
        self.assertEqual(unifier.as_number("1 234,5"), 1234.5)
        self.assertIsNone(unifier.as_number(True))
        self.assertIsNone(unifier.as_number(float("nan")))
        self.assertIsNone(unifier.as_number("=SUM(A1:A2)"))

    def test_detail_article_reconciliation_preserves_sizes(self):
        products, _ = parse_sheets(*invented_sheets())
        articles = {p.article for p in products}
        self.assertIn("MJMDEMO0001B", articles)
        self.assertNotIn("MJMDEMO0001A", articles)
        self.assertEqual(next(p for p in products if p.article == "MJMDEMO0001B").sizes,
                         {"XS": 3, "S": 5, "M": 7, "L": 11})
        self.assertEqual(next(p for p in products if p.article == "MJMDEMO0002P").sizes, {"L": 17})

    def test_total_rows_never_inherit_an_article(self):
        sheet = unifier.SheetData("Renamed pivot", [
            ["Название модели", "Цвет", "№ артикула", "S", "M", "Общее"],
            ["Учебная куртка", "Серый", "MJMDEMO0200", 7, 11, 18],
            ["Учебная куртка Total", "", "", 7, 11, 18],
            ["Grand Total", "", "", 7, 11, 18],
        ])
        products, _ = parse_sheets(sheet)
        self.assertEqual(sum(float(p.total or 0) for p in products), 18)
        self.assertTrue(unifier.is_service_or_total_row(["Учебный товар Total", "", 18]))
        self.assertFalse(any("TOTAL" in p.name.upper() for p in products))

    def test_photo_summary_selected_by_headers_not_name(self):
        weak = unifier.SheetData("Summary", [
            ["Название модели", "Цвет", "№ артикула", "S", "M", "Общее"],
            ["Учебная куртка", "Серый", "MJMDEMO0300", 2, 3, 5],
        ])
        preferred = unifier.SheetData("Arbitrary sender name", [
            ["№", "Фото модели", "Название модели", "№ артикула", "Цвет", "Цена", "S", "M", "Общее"],
            [1, "", "Учебная куртка", "MJMDEMO0300", "Серый", 2.25, 2, 3, 5],
        ])
        products, candidate = parse_sheets(weak, preferred)
        self.assertEqual(candidate.sheet.name, "Arbitrary sender name")
        self.assertEqual(sum(float(p.total or 0) for p in products), 5)

    def test_unrelated_articles_are_not_reconciled(self):
        first = unifier.SheetData("Summary", [
            ["Наименование", "Артикул", "Цвет", "S", "Общее"],
            ["Учебный шарф", "MJMDEMOAAAA", "Синий", 9, 9],
        ])
        other = unifier.SheetData("Detail", [
            ["Наименование", "Артикул", "Цвет", "S", "Общее"],
            ["Совершенно другой демонстрационный товар", "MJMDEMOBBBB", "Синий", 9, 9],
        ])
        products, _ = parse_sheets(first, other)
        self.assertEqual(products[0].article, "MJMDEMOAAAA")

    def test_merge_and_long_output_keep_canonical_size_order(self):
        products = unifier.merge_products([
            unifier.ProductRow("Demo", "MJMDEMO0400", "Blue", sizes={"S": 3, "M": 5}, total=8),
            unifier.ProductRow("demo", "mjmdemo0400", "blue", sizes={"S": 4, "M": 7}, total=11),
        ])
        self.assertEqual(len(products), 1)
        self.assertEqual(products[0].sizes, {"S": 7, "M": 12})
        self.assertEqual(products[0].total, 19)
        rows = unifier.expanded_output_rows([
            unifier.ProductRow("Demo", "MJMDEMO0500", sizes={"2XL": 3, "XS": 5, "S": 4})
        ])
        self.assertEqual([r[4] for r in rows], ["XS", "S", "2XL"])
        self.assertEqual([r[5] for r in rows], [5, 4, 3])

    def test_towel_dimensions_do_not_become_apparel_variations(self):
        rows = unifier.expanded_output_rows([
            unifier.ProductRow("Demo towel", "MJMDEMO0600", declared_size="30×40", total=9)
        ])
        self.assertEqual(rows[0][4], "")
        self.assertEqual(rows[0][5], 9)

    def test_output_cannot_overwrite_source_and_macro_output_is_xlsx(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "synthetic.xlsx"
            with self.assertRaisesRegex(ValueError, "исходный файл"):
                unifier.output_path_for(source, ".")
            self.assertEqual(unifier.output_path_for(source.with_suffix(".xlsm"), "Result").suffix, ".xlsx")
        with self.assertRaises(ValueError):
            unifier.read_workbook(Path("synthetic.csv"))

    def test_temporary_workbook_roundtrip_preserves_source(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "synthetic.xlsx"
            book = Workbook()
            book.active.title = "Cover"
            book.active.append(["Invented fixture; not an invoice"])
            sheet = book.create_sheet("Renamed product grid")
            sheet.append(["Фото модели", "Название модели", "№ артикула", "Цвет", "S", "M", "Общее"])
            sheet.append(["", "Учебная футболка", "MJMROUND0001", "Индиго", 5, 7, 12])
            sheet.append(["", "Учебное полотенце 30x40", "MJMROUND0002", "Белый", 0, 0, 3])
            sheet.append(["", "Учебная футболка Total", "", "", 5, 7, 12])
            sheet.append(["Grand Total", "", "", "", 5, 7, 15])
            book.save(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            output, candidate, count = unifier.convert_file(source)
            self.assertEqual(candidate.sheet.name, "Renamed product grid")
            self.assertEqual(count, 3)
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            self.assertNotEqual(output.resolve(), source.resolve())
            saved = load_workbook(output, read_only=True, data_only=True)
            try:
                rows = list(saved.active.values)
            finally:
                saved.close()
            self.assertEqual(list(rows[0]), HEADERS)
            self.assertEqual(sum(row[5] for row in rows[1:]), 15)
            self.assertEqual({(r[2], r[4] or "", r[5]) for r in rows[1:]},
                             {("MJMROUND0001", "S", 5), ("MJMROUND0001", "M", 7), ("MJMROUND0002", "", 3)})

    def test_written_article_remains_literal_text(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "synthetic-output.xlsx"
            count = unifier.build_output([unifier.ProductRow("Demo", "000123", total=4)], output, Path("synthetic-source.xlsx"))
            self.assertEqual(count, 1)
            saved = load_workbook(output, read_only=True, data_only=False)
            try:
                self.assertEqual(saved.active["C2"].value, "000123")
                self.assertEqual(saved.active["C2"].data_type, "s")
            finally:
                saved.close()

    def test_optional_private_fixture_structural_smoke(self):
        supplied = os.environ.get("EXCEL_UNIFIER_PRIVATE_FIXTURE")
        if not supplied:
            self.skipTest("optional private fixture not explicitly supplied")
        source = Path(supplied).expanduser().resolve(strict=True)
        products, _ = unifier.extract_products(source)
        self.assertTrue(products, "private fixture had no recognizable product table")
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "private-check.xlsx"
            count = unifier.build_output(products, output, Path("private-fixture.xlsx"))
            saved = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(list(next(saved.active.values)), HEADERS)
                self.assertEqual(count, saved.active.max_row - 1)
            finally:
                saved.close()


if __name__ == "__main__":
    unittest.main()
